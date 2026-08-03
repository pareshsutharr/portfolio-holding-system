"""Generate Stock_Style_Classification_Methodology_v1.docx.

Documentation aid only -- does not touch stock_style.py, stock_style_scoring.py,
stock_style_universe.py, stock_style_market.py, or style_config.json.

Updates from the original version:
  - Low Volatility and Liquidity sections removed entirely -- both are
    already calculated at the portfolio level in the Risk-O-Meter, so they
    are not duplicated here.
  - Momentum's Relative Strength formula corrected to match the code: a
    ratio, not a subtraction.
  - Value's P/E, P/B, EV/EBITDA and Dividend Yield are now benchmarked
    against the stock's NSE sectoral index, not an unspecified basis.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Stock_Style_Classification_Methodology_v1.docx"


def _data_required(document: Document, items: list[str]) -> None:
    document.add_paragraph("Data Required:")
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def _calculation(document: Document, formula: str, note: str) -> None:
    document.add_paragraph(f"{formula}\n{note}")


def generate() -> Path:
    document = Document()
    document.add_heading("Stock Style Classification Methodology (Version 1)", level=1)
    document.add_paragraph(
        "Objective: Define the data and calculations required to classify "
        "stocks into Growth, Value, Momentum and Quality styles, and to "
        "separately report a Size classification. Volatility and Liquidity "
        "risk are already calculated at the portfolio level in the "
        "Risk-O-Meter and are intentionally not recalculated here."
    )

    document.add_heading("Growth", level=2)
    _data_required(document, [
        "Revenue (5Y)", "EPS (5Y)", "Net Profit (5Y)", "ROE (5Y)", "ROCE (5Y)",
    ])
    document.add_paragraph("Calculations:")
    _calculation(document, "Revenue CAGR: ((Ending Revenue / Beginning Revenue)^(1/Years))-1", "Measures annualized revenue growth.")
    _calculation(document, "EPS CAGR: ((Ending EPS / Beginning EPS)^(1/Years))-1", "Measures annualized EPS growth.")
    _calculation(document, "Net Profit CAGR: ((Ending Profit / Beginning Profit)^(1/Years))-1", "Measures annualized profit growth.")
    _calculation(document, "Average ROE: Average of last 5 yearly ROE values", "Higher is better.")
    _calculation(document, "Average ROCE: Average of last 5 yearly ROCE values", "Higher is better.")

    document.add_heading("Value", level=2)
    _data_required(document, [
        "Latest P/E", "Latest P/B", "Latest EV/EBITDA", "Latest Dividend Yield",
        "The stock's NSE sectoral index P/E, P/B, EV/EBITDA and Dividend Yield",
    ])
    document.add_paragraph("Calculations:")
    _calculation(document, "P/E: Use latest value, compared against the stock's NSE sectoral index P/E.", "Lower relative to the sector benchmark is generally better.")
    _calculation(document, "P/B: Use latest value, compared against the stock's NSE sectoral index P/B.", "Lower relative to the sector benchmark is generally better.")
    _calculation(document, "EV/EBITDA: Use latest value, compared against the stock's NSE sectoral index EV/EBITDA.", "Lower relative to the sector benchmark is generally better.")
    _calculation(document, "Dividend Yield: Use latest value, compared against the stock's NSE sectoral index Dividend Yield.", "Higher relative to the sector benchmark is generally better.")
    document.add_paragraph(
        "Note: benchmarking against the stock's own NSE sectoral index (rather "
        "than the whole market) avoids penalizing naturally high-multiple "
        "sectors (e.g. IT) or over-rewarding naturally low-multiple sectors "
        "(e.g. PSUs, cyclicals)."
    )

    document.add_heading("Momentum", level=2)
    _data_required(document, ["1 year daily stock close", "1 year daily Nifty50 close"])
    document.add_paragraph("Calculations:")
    _calculation(document, "6 Month Return: (Current Price / Price 6 Months Ago)-1", "Example: 250/200=1.25;1.25-1=0.25=25%")
    _calculation(document, "12 Month Return: (Current Price / Price 12 Months Ago)-1", "Same approach over 12 months.")
    _calculation(
        document,
        "Relative Strength: (1 + 12M Stock Return) / (1 + 12M Nifty Return)",
        "A ratio above 1 means outperformance versus Nifty over the trailing 12 months.",
    )
    _calculation(document, "200 DMA: Current Price > 200-day Moving Average", "Above=positive trend.")

    document.add_heading("Quality", level=2)
    _data_required(document, ["ROE", "ROCE", "Debt/Equity", "Interest Coverage"])
    document.add_paragraph("Calculations:")
    _calculation(document, "ROE: Latest year", "Higher better.")
    _calculation(document, "ROCE: Latest year", "Higher better.")
    _calculation(document, "Debt/Equity: Latest", "Lower better.")
    _calculation(document, "Interest Coverage: Latest", "Higher better.")

    document.add_heading("Size", level=2)
    _data_required(document, ["Latest Market Cap"])
    document.add_paragraph("Calculations:")
    _calculation(document, "Classification: Large/Mid/Small/Micro", "Based on NSE/AMFI definitions.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(generate())
