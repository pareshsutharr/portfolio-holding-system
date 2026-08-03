"""Generate a plain-language, worked-example guide to calculating the
Portfolio Risk-O-Meter score by hand.

This is a documentation aid only -- it does not touch riskometer.py or
riskometer_config.json, and it performs no live calculation of its own.
Every number in the worked example below was computed by hand from the
same formulas and thresholds already defined in riskometer_config.json,
so it can be followed with a calculator or a spreadsheet.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Riskometer_Manual_Calculation_Guide_v1.docx"


def _add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)


def generate() -> Path:
    document = Document()
    document.add_heading(
        "How to Calculate the Portfolio Risk-O-Meter Score (Step-by-Step Guide)", level=1
    )
    document.add_paragraph(
        "This guide walks through the calculation by hand, using one worked "
        "example portfolio of 6 stocks. Every number below can be reproduced "
        "in a calculator or spreadsheet. All scores are on a 0-100 scale "
        "where 100 = highest risk."
    )

    document.add_heading("Step 0: The Example Portfolio", level=1)
    document.add_paragraph(
        "Assume you already know each stock's weight, sector, market-cap "
        "category, latest fundamentals, recent trading activity, historical "
        "volatility figures and beta (these come from your holdings file, "
        "company financials and market data -- this guide starts from there)."
    )
    _add_table(
        document,
        ["Stock", "Weight", "Sector", "Cap", "ROE%", "ROCE%", "D/E", "Int. Cover"],
        [
            ["A", "24%", "IT", "Large", "22", "20", "0.3", "15x"],
            ["B", "20%", "Banking (financial)", "Large", "15", "13", "1.2", "N/A"],
            ["C", "16%", "Auto", "Mid", "12", "11", "0.9", "6x"],
            ["D", "14%", "Pharma", "Large", "17", "16", "0.4", "12x"],
            ["E", "13%", "FMCG", "Mid", "19", "17", "0.2", "20x"],
            ["F", "13%", "Metals", "Small", "6", "5", "1.9", "2.5x"],
        ],
    )
    _add_table(
        document,
        ["Stock", "Avg Volume ('000)", "Avg Turnover (Cr)", "Ann. Volatility", "Max Drawdown", "Downside Vol.", "Beta"],
        [
            ["A", "600", "50", "30%", "25%", "20%", "1.05"],
            ["B", "900", "70", "28%", "22%", "18%", "0.95"],
            ["C", "200", "15", "40%", "33%", "28%", "1.20"],
            ["D", "350", "25", "25%", "18%", "15%", "0.70"],
            ["E", "150", "8", "22%", "16%", "13%", "0.60"],
            ["F", "30", "1.5", "60%", "48%", "42%", "1.55"],
        ],
    )
    document.add_paragraph(
        "Note: bank B has no meaningful Interest Cover ratio, so it is "
        "excluded from that one metric only (this is the standard rule for "
        "banks/NBFCs/insurers)."
    )

    document.add_heading("Master Threshold & Weight Reference", level=1)
    document.add_paragraph(
        "Keep this table open while you work through every step below -- "
        "every threshold used in the guide comes from here."
    )
    _add_table(
        document,
        ["Module", "Weight", "Metric", "Threshold"],
        [
            ["Concentration", "15%", "Largest Holding", "High = 40%"],
            ["", "", "Top 3 Holdings", "High = 70%"],
            ["", "", "Top 5 Holdings", "High = 85%"],
            ["", "", "Holding HHI", "High = 0.25"],
            ["Sector", "15%", "Largest Sector", "High = 50%"],
            ["", "", "Sector HHI", "High = 0.30"],
            ["", "", "Sector Count", "Good = 10 sectors"],
            ["Market Cap", "10%", "Large / Mid / Small / Micro / Unknown", "15 / 45 / 75 / 95 / 80"],
            ["Fundamental Quality", "15%", "ROE", "Good = 20%"],
            ["", "", "ROCE", "Good = 20%"],
            ["", "", "Debt/Equity", "High = 2.0"],
            ["", "", "Interest Cover", "Good = 8x"],
            ["Liquidity (Large Cap)", "10%", "Avg Volume ('000) / Turnover (Cr)", "Low = 50 / 5, High = 2000 / 200"],
            ["Liquidity (Mid Cap)", "", "Avg Volume ('000) / Turnover (Cr)", "Low = 20 / 2, High = 800 / 80"],
            ["Liquidity (Small Cap)", "", "Avg Volume ('000) / Turnover (Cr)", "Low = 5 / 0.5, High = 300 / 30"],
            ["Volatility", "20%", "Annualized Volatility", "High = 60%"],
            ["", "", "Maximum Drawdown", "High = 50%"],
            ["", "", "Downside Volatility", "High = 45%"],
            ["Beta", "15%", "Beta (benchmark by cap: Nifty 50 / Nifty Midcap 150 / Nifty 500)", "Low = 0.60, High = 1.60"],
        ],
    )

    document.add_heading("Step 1: Concentration Risk (Weight 15%)", level=1)
    document.add_paragraph(
        "1. Sort weights largest to smallest: 24, 20, 16, 14, 13, 13."
    )
    document.add_paragraph("2. Largest holding = 24%.")
    document.add_paragraph("3. Top 3 = 24 + 20 + 16 = 60%.")
    document.add_paragraph("4. Top 5 = 24 + 20 + 16 + 14 + 13 = 87%.")
    document.add_paragraph(
        "5. HHI = sum of each weight squared (as a fraction, not a %): "
        "0.24^2 + 0.20^2 + 0.16^2 + 0.14^2 + 0.13^2 + 0.13^2 = 0.1766."
    )
    document.add_paragraph(
        "6. Turn each into a risk score with: Risk = 100 x min(value, threshold) / threshold. "
        "If value exceeds the threshold it is simply capped at 100 (maximum risk)."
    )
    _add_table(
        document,
        ["Metric", "Value", "Threshold", "Calculation", "Risk Score"],
        [
            ["Largest Holding", "24%", "40%", "100 x 24/40", "60.00"],
            ["Top 3", "60%", "70%", "100 x 60/70", "85.71"],
            ["Top 5", "87%", "85%", "100 x min(87,85)/85 -> capped", "100.00"],
            ["HHI", "0.1766", "0.25", "100 x 0.1766/0.25", "70.64"],
        ],
    )
    document.add_paragraph(
        "7. Concentration Risk = average of the four risk scores = "
        "(60.00 + 85.71 + 100.00 + 70.64) / 4 = 79.09."
    )

    document.add_heading("Step 2: Sector Risk (Weight 15%)", level=1)
    document.add_paragraph(
        "Each stock in this example is in a different sector, so sector "
        "weights equal stock weights: IT 24%, Banking 20%, Auto 16%, "
        "Pharma 14%, FMCG 13%, Metals 13% -- 6 sectors in total."
    )
    _add_table(
        document,
        ["Metric", "Value", "Threshold", "Calculation", "Risk Score"],
        [
            ["Largest Sector (IT)", "24%", "High = 50%", "100 x 24/50", "48.00"],
            ["Sector HHI", "0.1766", "High = 0.30", "100 x 0.1766/0.30", "58.87"],
            [
                "Sector Count",
                "6 sectors",
                "Good = 10",
                "100 x (1 - min(6,10)/10)",
                "40.00",
            ],
        ],
    )
    document.add_paragraph(
        "Sector Count uses the opposite direction (more sectors = lower "
        "risk), so it is: Risk = 100 x (1 - value/good), floored at 0."
    )
    document.add_paragraph(
        "Sector Risk = average of the three = (48.00 + 58.87 + 40.00) / 3 = 48.96."
    )

    document.add_heading("Step 3: Market Cap Risk (Weight 10%)", level=1)
    document.add_paragraph(
        "Group holdings by SEBI market-cap category and multiply each "
        "category's total weight by its configured risk score, then add "
        "them up."
    )
    _add_table(
        document,
        ["Category", "Combined Weight", "Configured Score", "Contribution"],
        [
            ["Large Cap (A, B, D)", "58%", "15", "0.58 x 15 = 8.70"],
            ["Mid Cap (C, E)", "29%", "45", "0.29 x 45 = 13.05"],
            ["Small Cap (F)", "13%", "75", "0.13 x 75 = 9.75"],
        ],
    )
    document.add_paragraph("Market Cap Risk = 8.70 + 13.05 + 9.75 = 31.50.")

    document.add_heading("Step 4: Fundamental Quality Risk (Weight 15%)", level=1)
    document.add_paragraph(
        "For each stock: ROE and ROCE use Risk = 100 x (1 - min(value, 20) "
        "/ 20). Debt/Equity uses Risk = 100 x min(value, 2) / 2. Interest "
        "Cover (non-financial stocks only) uses Risk = 100 x (1 - min(value, "
        "8) / 8). Average the applicable component risks to get each "
        "stock's Quality Risk."
    )
    _add_table(
        document,
        ["Stock", "ROE Risk", "ROCE Risk", "D/E Risk", "Int. Cover Risk", "Stock Quality Risk"],
        [
            ["A", "0.00", "0.00", "15.00", "0.00", "3.75"],
            ["B (financial)", "25.00", "35.00", "60.00", "N/A", "40.00"],
            ["C", "40.00", "45.00", "45.00", "25.00", "38.75"],
            ["D", "15.00", "20.00", "20.00", "0.00", "13.75"],
            ["E", "5.00", "15.00", "10.00", "0.00", "7.50"],
            ["F", "70.00", "75.00", "95.00", "68.75", "77.19"],
        ],
    )
    document.add_paragraph(
        "Portfolio Quality Risk = weighted average of the Stock Quality Risk "
        "column, using portfolio weights: (0.24 x 3.75) + (0.20 x 40.00) + "
        "(0.16 x 38.75) + (0.14 x 13.75) + (0.13 x 7.50) + (0.13 x 77.19) = "
        "28.03."
    )
    document.add_paragraph(
        "Note: ETFs, if any, are skipped entirely in this module and do not "
        "enter the weighted average."
    )

    document.add_heading("Step 5: Liquidity Risk (Weight 10%)", level=1)
    document.add_paragraph(
        "Volume and Turnover use a logarithmic scale because trading "
        "activity spans a very wide range. For a value between Low and "
        "High: Risk = 100 x (1 - (ln(value) - ln(Low)) / (ln(High) - "
        "ln(Low))). At or below Low, Risk = 100. At or above High, Risk = 0."
    )
    document.add_paragraph(
        "Unlike every other module, Liquidity uses a DIFFERENT Low/High pair "
        "depending on each stock's own market-cap category (see the Master "
        "Threshold table): Large Cap stocks are judged against a much higher "
        "bar than Small Cap stocks, because what counts as 'thinly traded' "
        "is naturally different for a large-cap versus a small-cap company."
    )
    _add_table(
        document,
        ["Stock", "Cap", "Volume Risk", "Turnover Risk", "Stock Liquidity Risk"],
        [
            ["A", "Large", "32.64", "37.58", "35.11"],
            ["B", "Large", "21.65", "28.46", "25.05"],
            ["C", "Mid", "37.58", "45.38", "41.48"],
            ["D", "Large", "47.25", "56.37", "51.81"],
            ["E", "Mid", "45.38", "62.42", "53.90"],
            ["F", "Small", "56.24", "73.17", "64.70"],
        ],
    )
    document.add_paragraph(
        "Portfolio Liquidity Risk = weighted average of Stock Liquidity Risk "
        "= (0.24 x 35.11) + (0.20 x 25.05) + (0.16 x 41.48) + (0.14 x 51.81) "
        "+ (0.13 x 53.90) + (0.13 x 64.70) = 42.75."
    )
    document.add_paragraph(
        "Notice Large Cap stocks A, B and D now score noticeably worse than "
        "under a single universal threshold, and Small Cap stock F scores "
        "noticeably better -- because each is now being judged against a "
        "bar appropriate to its own size, not one bar for everyone."
    )

    document.add_heading("Step 6: Volatility Risk (Weight 20%)", level=1)
    document.add_paragraph(
        "Each of the three volatility metrics uses: Risk = 100 x min(value, "
        "threshold) / threshold. Average the three for each stock."
    )
    _add_table(
        document,
        ["Stock", "Volatility Risk", "Drawdown Risk", "Downside Risk", "Stock Volatility Risk"],
        [
            ["A", "50.00", "50.00", "44.44", "48.15"],
            ["B", "46.67", "44.00", "40.00", "43.56"],
            ["C", "66.67", "66.00", "62.22", "64.96"],
            ["D", "41.67", "36.00", "33.33", "37.00"],
            ["E", "36.67", "32.00", "28.89", "32.52"],
            ["F", "100.00", "96.00", "93.33", "96.44"],
        ],
    )
    document.add_paragraph(
        "Portfolio Volatility Risk = weighted average of Stock Volatility "
        "Risk = (0.24 x 48.15) + (0.20 x 43.56) + (0.16 x 64.96) + (0.14 x "
        "37.00) + (0.13 x 32.52) + (0.13 x 96.44) = 52.61."
    )

    document.add_heading("Step 7: Beta Risk (Weight 15%)", level=1)
    document.add_paragraph(
        "Risk = 100 x (Beta - 0.60) / (1.60 - 0.60), clipped to the 0-100 "
        "range. Since the range is exactly 1.00 wide here, this simplifies "
        "to Risk = 100 x (Beta - 0.60)."
    )
    document.add_paragraph(
        "Each stock's Beta is measured against the benchmark that matches "
        "its own market-cap category, not against Nifty 50 for every stock: "
        "Large Cap stocks use Nifty 50, Mid Cap stocks use Nifty Midcap 150, "
        "and Small Cap stocks use Nifty 500. The risk-scoring formula and "
        "the 0.60-1.60 band stay the same either way -- only the benchmark "
        "the Beta itself was measured against changes."
    )
    _add_table(
        document,
        ["Stock", "Cap", "Benchmark Used", "Beta", "Calculation", "Risk Score"],
        [
            ["A", "Large", "Nifty 50", "1.05", "100 x (1.05 - 0.60)", "45.00"],
            ["B", "Large", "Nifty 50", "0.95", "100 x (0.95 - 0.60)", "35.00"],
            ["C", "Mid", "Nifty Midcap 150", "1.20", "100 x (1.20 - 0.60)", "60.00"],
            ["D", "Large", "Nifty 50", "0.70", "100 x (0.70 - 0.60)", "10.00"],
            ["E", "Mid", "Nifty Midcap 150", "0.60", "100 x (0.60 - 0.60)", "0.00"],
            ["F", "Small", "Nifty 500", "1.55", "100 x (1.55 - 0.60)", "95.00"],
        ],
    )
    document.add_paragraph(
        "Note: the Beta values above are given, illustrative inputs for "
        "this example (as with every other input in this guide), so they "
        "are unchanged from the earlier version of this walkthrough. In the "
        "actual system, recomputing C, E and F's Beta against their own "
        "cap-appropriate benchmark instead of Nifty 50 would generally "
        "produce different numeric Beta values, since a stock's covariance "
        "with a mid-cap or small-cap index is not the same as its "
        "covariance with Nifty 50."
    )
    document.add_paragraph(
        "Portfolio Beta Risk = weighted average = (0.24 x 45.00) + (0.20 x "
        "35.00) + (0.16 x 60.00) + (0.14 x 10.00) + (0.13 x 0.00) + (0.13 x "
        "95.00) = 41.15."
    )

    document.add_heading("Step 8: Combine Into the Overall Score", level=1)
    document.add_paragraph(
        "Multiply each module's risk score by its weight and add them all "
        "together."
    )
    _add_table(
        document,
        ["Module", "Score", "Weight", "Contribution"],
        [
            ["Concentration", "79.09", "15%", "11.86"],
            ["Sector", "48.96", "15%", "7.34"],
            ["Market Cap", "31.50", "10%", "3.15"],
            ["Fundamental Quality", "28.03", "15%", "4.21"],
            ["Liquidity", "42.75", "10%", "4.28"],
            ["Volatility", "52.61", "20%", "10.52"],
            ["Beta", "41.15", "15%", "6.17"],
            ["Overall", "-", "100%", "47.53"],
        ],
    )
    document.add_paragraph(
        "Overall Score = 11.86 + 7.34 + 3.15 + 4.21 + 4.28 + 10.52 + 6.17 = 47.53."
    )

    document.add_heading("Step 9: Read Off the Rating", level=1)
    _add_table(
        document,
        ["Score Range", "Rating"],
        [
            ["0-20", "Very Low"],
            ["20-40", "Low"],
            ["40-60", "Moderate"],
            ["60-80", "High"],
            ["80-100", "Very High"],
        ],
    )
    document.add_paragraph(
        "47.53 falls between 40 and 60, so this example portfolio is rated "
        "Moderate risk."
    )

    document.add_heading("Summary", level=1)
    document.add_paragraph(
        "Every module follows the same three-step pattern: (1) get the raw "
        "metric per stock, (2) convert it to a 0-100 risk score against its "
        "threshold, (3) take the portfolio-weighted average across stocks "
        "(or, for Concentration/Sector/Market Cap, across the portfolio "
        "structure directly). The seven module scores are then combined "
        "with their fixed weights into one overall score, which is mapped "
        "to a rating band."
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(generate())
