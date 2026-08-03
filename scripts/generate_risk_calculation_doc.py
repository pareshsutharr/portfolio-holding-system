"""Generate a transparent stock-by-stock Risk-O-Meter calculation document."""

from __future__ import annotations

import json
import math
import re
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from sqlalchemy import bindparam, text

from market_etl.config import Settings
from market_etl.database import build_engine
from portfolio_returns import load_all_benchmark_series
from riskometer import Riskometer, _clip, _direct_bad, _inverse_good, _log_inverse
from scripts.run_riskometer import read_holdings


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "portfolio_risk_calculations.docx"


def _fmt(value: Any, decimals: int = 2, suffix: str = "") -> str:
    if value is None or pd.isna(value):
        return "N/A"
    return f"{float(value):,.{decimals}f}{suffix}"


def _shade(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def _set_cell_text(cell: Any, value: Any, *, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(value))
    run.bold = bold
    run.font.size = Pt(8)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _add_table(document: Document, headers: list[str], rows: list[list[Any]]) -> Any:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        _set_cell_text(table.rows[0].cells[index], header, bold=True)
        _shade(table.rows[0].cells[index], "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            _set_cell_text(cells[index], value)
    return table


def _component_formula(label: str, value: float, threshold: float, score: float) -> str:
    return (
        f"{label}: min(max({value:.4f}, 0), {threshold:.4f}) "
        f"/ {threshold:.4f} x 100 = {score:.2f}"
    )


def _load_data() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, dict[str, Any]]:
    config = json.loads((ROOT / "riskometer_config.json").read_text(encoding="utf-8"))
    holdings = read_holdings(ROOT / "test_files" / "holdings.xlsx")
    holdings["isin"] = holdings["isin"].astype(str).str.strip().str.upper()
    holdings["value"] = pd.to_numeric(holdings["value"], errors="coerce")
    holdings = holdings.dropna(subset=["isin", "value"])
    holdings = (
        holdings[holdings["value"] > 0]
        .groupby("isin", as_index=False)
        .agg(security_name=("security_name", "first"), value=("value", "sum"))
    )
    holdings["weight"] = holdings["value"] / holdings["value"].sum()
    isins = holdings["isin"].tolist()

    engine = build_engine(Settings.from_env().database_url)
    with engine.connect() as connection:
        master = pd.read_sql(
            text(
                "SELECT isin, company_name, sector, industry "
                "FROM company_master WHERE isin IN :isins"
            ).bindparams(bindparam("isins", expanding=True)),
            connection,
            params={"isins": isins},
        )
        fundamentals = pd.read_sql(
            text(
                "SELECT DISTINCT ON (isin) isin, financial_year, roe, roce, "
                "interest_cover, debt_equity FROM fundamentals "
                "WHERE isin IN :isins ORDER BY isin, financial_year DESC"
            ).bindparams(bindparam("isins", expanding=True)),
            connection,
            params={"isins": isins},
        )
        market = pd.read_sql(
            text(
                "SELECT isin, trade_date, close_price, volume, turnover "
                "FROM daily_market_data WHERE isin IN :isins ORDER BY isin, trade_date"
            ).bindparams(bindparam("isins", expanding=True)),
            connection,
            params={"isins": isins},
        )

    benchmark_prices = load_all_benchmark_series()
    benchmark = {
        name: series.pct_change().dropna()
        for name, series in benchmark_prices.items()
    }

    cap = pd.read_excel(ROOT / "test_files" / "MCap.xlsx", header=1)[
        ["ISIN", "Categorization as per SEBI Circular dated Oct 6, 2017"]
    ].rename(
        columns={
            "ISIN": "isin",
            "Categorization as per SEBI Circular dated Oct 6, 2017": "cap_category",
        }
    )
    cap["isin"] = cap["isin"].astype(str).str.strip().str.upper()
    cap = cap.drop_duplicates("isin")
    sector_mapping = json.loads(
        (ROOT / "sector_mapping.json").read_text(encoding="utf-8")
    )

    portfolio = (
        holdings.merge(master, on="isin", how="left")
        .merge(cap, on="isin", how="left")
        .merge(fundamentals, on="isin", how="left")
    )
    portfolio["nse_sector"] = (
        portfolio["sector"].map(sector_mapping).fillna(portfolio["sector"]).fillna("Unmapped")
    )
    portfolio["cap_category"] = portfolio["cap_category"].fillna("Unknown")
    labels = (
        portfolio[["security_name", "sector", "nse_sector"]]
        .fillna("")
        .astype(str)
        .agg(" ".join, axis=1)
        .str.upper()
    )
    portfolio["is_etf"] = labels.str.contains(
        r"\bETF\b|EXCHANGE TRADED|GOLD\s*BEES|NETF(?:GOLD|SILVER)", regex=True
    )
    portfolio["is_financial"] = labels.str.contains(
        r"\bBANK\b|FINANCIAL SERVICES|\bFINANCE\b|NBFC|INSURANCE", regex=True
    )
    return portfolio, market, benchmark, config


def _calculate_stock_metrics(
    portfolio: pd.DataFrame,
    market: pd.DataFrame,
    benchmark: pd.DataFrame,
    config: dict[str, Any],
) -> pd.DataFrame:
    result = portfolio.copy()
    quality_config = config["quality"]

    def quality_components(row: pd.Series) -> pd.Series:
        if row["is_etf"]:
            return pd.Series(
                {
                    "roe_risk": np.nan,
                    "roce_risk": np.nan,
                    "debt_equity_risk": np.nan,
                    "interest_cover_risk": np.nan,
                    "quality_score": np.nan,
                    "quality_status": "Excluded: ETF fundamentals are Not Applicable",
                }
            )
        required = ["roe", "roce", "debt_equity"]
        if not row["is_financial"]:
            required.append("interest_cover")
        if any(pd.isna(row[column]) for column in required):
            return pd.Series(
                {
                    "roe_risk": np.nan,
                    "roce_risk": np.nan,
                    "debt_equity_risk": np.nan,
                    "interest_cover_risk": np.nan,
                    "quality_score": np.nan,
                    "quality_status": "Not scored: an applicable fundamental is missing",
                }
            )
        roe_risk = _inverse_good(float(row["roe"]), quality_config["roe_good"])
        roce_risk = _inverse_good(float(row["roce"]), quality_config["roce_good"])
        debt_risk = _direct_bad(
            float(row["debt_equity"]), quality_config["debt_equity_high"]
        )
        components = [roe_risk, roce_risk, debt_risk]
        interest_risk = np.nan
        if not row["is_financial"]:
            interest_risk = _inverse_good(
                float(row["interest_cover"]), quality_config["interest_cover_good"]
            )
            components.append(interest_risk)
        return pd.Series(
            {
                "roe_risk": roe_risk,
                "roce_risk": roce_risk,
                "debt_equity_risk": debt_risk,
                "interest_cover_risk": interest_risk,
                "quality_score": float(np.mean(components)),
                "quality_status": (
                    "Scored on ROE, ROCE and Debt/Equity; Interest Cover N/A"
                    if row["is_financial"]
                    else "Scored on all four applicable fundamentals"
                ),
            }
        )

    result = pd.concat([result, result.apply(quality_components, axis=1)], axis=1)

    recent = (
        market.sort_values("trade_date")
        .groupby("isin")
        .tail(config["history"]["liquidity_days"])
    )
    liquidity = (
        recent.groupby("isin")
        .agg(
            liquidity_observations=("trade_date", "nunique"),
            average_volume=("volume", "mean"),
            average_turnover=("turnover", "mean"),
        )
        .reset_index()
    )
    liquidity = liquidity.merge(
        result[["isin", "cap_category"]], on="isin", how="left"
    )
    liquidity_tiers = config["liquidity"]
    fallback_tier = liquidity_tiers[config["liquidity_fallback_tier"]]

    def _liquidity_tier(cap_category: Any) -> dict[str, float]:
        return liquidity_tiers.get(cap_category, fallback_tier)

    liquidity["volume_risk"] = liquidity.apply(
        lambda row: _log_inverse(
            float(row["average_volume"]),
            _liquidity_tier(row["cap_category"])["volume_low_thousand"],
            _liquidity_tier(row["cap_category"])["volume_high_thousand"],
        )
        if pd.notna(row["average_volume"])
        else np.nan,
        axis=1,
    )
    liquidity["turnover_risk"] = liquidity.apply(
        lambda row: _log_inverse(
            float(row["average_turnover"]),
            _liquidity_tier(row["cap_category"])["turnover_low_crore"],
            _liquidity_tier(row["cap_category"])["turnover_high_crore"],
        )
        if pd.notna(row["average_turnover"])
        else np.nan,
        axis=1,
    )
    liquidity["liquidity_score"] = liquidity[
        ["volume_risk", "turnover_risk"]
    ].mean(axis=1)
    liquidity = liquidity.drop(columns=["cap_category"])

    return_rows: list[dict[str, Any]] = []
    for isin, group in market.groupby("isin"):
        prices = (
            group.dropna(subset=["close_price"])
            .sort_values("trade_date")
            .tail(config["history"]["trading_days"])
        )
        returns = (
            prices.set_index("trade_date")["close_price"]
            .astype(float)
            .pct_change()
            .dropna()
        )
        if len(returns) < config["history"]["minimum_return_observations"]:
            continue
        wealth = (1 + returns).cumprod()
        drawdown = wealth / wealth.cummax() - 1
        downside = returns[returns < 0]
        return_rows.append(
            {
                "isin": isin,
                "return_observations": len(returns),
                "annual_volatility": returns.std() * np.sqrt(252),
                "maximum_drawdown": abs(drawdown.min()),
                "downside_volatility": (
                    downside.std() * np.sqrt(252) if len(downside) > 1 else 0
                ),
                "returns": returns,
            }
        )
    returns_frame = pd.DataFrame(return_rows)
    volatility_config = config["volatility"]
    if not returns_frame.empty:
        returns_frame["annual_volatility_risk"] = returns_frame[
            "annual_volatility"
        ].map(
            lambda value: _direct_bad(
                float(value), volatility_config["annualized_volatility_high"]
            )
        )
        returns_frame["drawdown_risk"] = returns_frame["maximum_drawdown"].map(
            lambda value: _direct_bad(
                float(value), volatility_config["maximum_drawdown_high"]
            )
        )
        returns_frame["downside_risk"] = returns_frame[
            "downside_volatility"
        ].map(
            lambda value: _direct_bad(
                float(value), volatility_config["downside_volatility_high"]
            )
        )
        returns_frame["volatility_score"] = returns_frame[
            ["annual_volatility_risk", "drawdown_risk", "downside_risk"]
        ].mean(axis=1)

    beta_cfg = config["beta"]
    benchmark_by_cap = beta_cfg["benchmark_by_cap"]
    fallback_benchmark = beta_cfg["benchmark_fallback"]
    cap_by_isin = result.set_index("isin")["cap_category"]

    beta_rows: list[dict[str, Any]] = []
    for row in returns_frame.itertuples():
        cap_category = cap_by_isin.get(row.isin, "Unknown")
        benchmark_name = benchmark_by_cap.get(cap_category, fallback_benchmark)
        benchmark_series = benchmark[benchmark_name]
        aligned = pd.concat(
            [
                row.returns.rename("stock"),
                benchmark_series.rename("benchmark"),
            ],
            axis=1,
        ).dropna()
        if (
            len(aligned) < config["history"]["minimum_return_observations"]
            or aligned["benchmark"].var() == 0
        ):
            continue
        beta = aligned["stock"].cov(aligned["benchmark"]) / aligned["benchmark"].var()
        low = beta_cfg["low"]
        high = beta_cfg["high"]
        beta_rows.append(
            {
                "isin": row.isin,
                "beta_observations": len(aligned),
                "beta": beta,
                "beta_score": _clip((float(beta) - low) / (high - low) * 100),
                "beta_benchmark": benchmark_name,
            }
        )
    beta_frame = pd.DataFrame(beta_rows)

    if "returns" in returns_frame:
        returns_frame = returns_frame.drop(columns=["returns"])
    result = result.merge(liquidity, on="isin", how="left")
    result = result.merge(returns_frame, on="isin", how="left")
    if not beta_frame.empty:
        result = result.merge(beta_frame, on="isin", how="left")
    else:
        result["beta_observations"] = np.nan
        result["beta"] = np.nan
        result["beta_score"] = np.nan
        result["beta_benchmark"] = np.nan
    return result.sort_values("weight", ascending=False).reset_index(drop=True)


def _add_document_header(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Portfolio Risk-O-Meter\nDetailed Stock-by-Stock Calculations")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(15, 23, 42)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        f"Generated {datetime.now():%d %B %Y, %H:%M} | "
        "Transparent input data, formulas, normalized scores and portfolio aggregation"
    )


def _add_methodology(
    document: Document,
    metrics: pd.DataFrame,
    analysis: dict[str, Any],
    config: dict[str, Any],
) -> None:
    document.add_heading("1. Executive Result and Rating Structure", level=1)
    document.add_paragraph(
        f"Total portfolio value analysed: Rs. {analysis['total_value']:,.2f}. "
        f"Final Risk-O-Meter score: {analysis['overall_score']:.2f}/100 "
        f"({analysis['overall_level']} Risk). A higher score means higher measured risk."
    )
    rows = []
    weights = config["module_weights"]
    for result in analysis["results"]:
        contribution = result.score * weights[result.name]
        rows.append(
            [
                result.name.replace("_", " ").title(),
                _fmt(result.score),
                f"{weights[result.name] * 100:.0f}%",
                _fmt(contribution),
                result.level,
                result.coverage,
            ]
        )
    rows.append(
        [
            "Overall",
            _fmt(analysis["overall_score"]),
            "100%",
            _fmt(analysis["overall_score"]),
            analysis["overall_level"],
            "Weighted module score",
        ]
    )
    _add_table(
        document,
        ["Module", "Score", "Weight", "Contribution", "Level", "Coverage"],
        rows,
    )
    document.add_paragraph(
        "Overall formula: Sum(Module score x Module weight) = "
        + " + ".join(
            f"{result.score:.2f} x {weights[result.name]:.2f}"
            for result in analysis["results"]
        )
        + f" = {analysis['overall_score']:.2f}."
    )

    document.add_heading("2. Rules Applied Before Stock-Level Calculations", level=1)
    for text_value in [
        "Portfolio weight = stock holding value / total portfolio value.",
        "ETFs remain in concentration, sector, market-cap, liquidity, volatility and beta modules, but are excluded from company Fundamental Quality scoring.",
        "Banks, NBFCs and other financial-service companies are scored on ROE, ROCE and Debt/Equity. Interest Coverage is marked Not Applicable.",
        "Non-financial operating companies require ROE, ROCE, Debt/Equity and Interest Coverage.",
        "Liquidity uses the latest 63 available trading days, against a Low/High volume and turnover band specific to the stock's own market-cap category (Large, Mid or Small Cap).",
        "Volatility uses up to 252 trading days and requires at least 60 return observations.",
        "Beta uses stock returns against a cap-appropriate benchmark on overlapping dates and requires at least 60 observations: Nifty 50 for Large Cap, Nifty Midcap 150 for Mid Cap, Nifty 500 for Small Cap.",
        "Module-level stock scores are portfolio-weighted over the holdings for which that module is calculable. Coverage is disclosed.",
    ]:
        document.add_paragraph(text_value, style="List Bullet")

    document.add_heading("3. Portfolio-Level Allocation Modules", level=1)
    concentration = next(r for r in analysis["results"] if r.name == "concentration")
    sector = next(r for r in analysis["results"] if r.name == "sector")
    market_cap = next(r for r in analysis["results"] if r.name == "market_cap")
    document.add_paragraph(
        "Concentration Risk formula: average of normalized Largest Holding, Top 3, "
        "Top 5 and holding HHI risks."
    )
    document.add_paragraph(
        f"Values: {concentration.values}. Result: {concentration.score:.2f} "
        f"({concentration.level})."
    )
    document.add_paragraph(
        "Sector Risk formula: average of normalized Largest Sector, Sector HHI and "
        "sector-count risks."
    )
    document.add_paragraph(
        f"Values: {sector.values}. Result: {sector.score:.2f} ({sector.level})."
    )
    document.add_paragraph(
        "Market-Cap Risk formula: sum(portfolio allocation in category x configured "
        "category score). Category scores are Large 15, Mid 45, Small 75, "
        "Micro 95 and Unknown 80."
    )
    document.add_paragraph(
        f"Values: {market_cap.values}. Result: {market_cap.score:.2f} "
        f"({market_cap.level})."
    )

    weights = metrics[["security_name", "value", "weight", "nse_sector", "cap_category"]]
    allocation_rows = [
        [
            row.security_name,
            f"Rs. {row.value:,.2f}",
            f"{row.weight * 100:.2f}%",
            row.nse_sector,
            row.cap_category,
            f"{row.weight**2:.6f}",
        ]
        for row in weights.itertuples()
    ]
    _add_table(
        document,
        ["Security", "Holding Value", "Weight", "NSE Sector", "Market Cap", "HHI Contribution"],
        allocation_rows,
    )


def _add_stock_section(
    document: Document,
    row: pd.Series,
    rank: int,
    metrics: pd.DataFrame,
    config: dict[str, Any],
) -> None:
    document.add_page_break()
    document.add_heading(f"{rank}. {row['security_name']}", level=1)
    _add_table(
        document,
        ["ISIN", "Holding Value", "Portfolio Weight", "Sector", "Market Cap"],
        [[
            row["isin"],
            f"Rs. {row['value']:,.2f}",
            f"{row['weight'] * 100:.4f}%",
            row["nse_sector"],
            row["cap_category"],
        ]],
    )

    document.add_heading("Allocation and concentration data", level=2)
    top3 = rank <= 3
    top5 = rank <= 5
    document.add_paragraph(
        f"Weight formula: Rs. {row['value']:,.2f} / "
        f"Rs. {metrics['value'].sum():,.2f} x 100 = {row['weight'] * 100:.4f}%. "
        f"Holding-HHI contribution = {row['weight']:.6f}² = "
        f"{row['weight'] ** 2:.6f}. Included in Top 3: {'Yes' if top3 else 'No'}; "
        f"included in Top 5: {'Yes' if top5 else 'No'}."
    )
    cap_score = config["market_cap_scores"].get(
        row["cap_category"], config["market_cap_scores"]["Unknown"]
    )
    document.add_paragraph(
        f"Market-cap contribution = {row['weight']:.6f} x {cap_score} = "
        f"{row['weight'] * cap_score:.4f} risk points."
    )

    document.add_heading("Fundamental Quality calculation", level=2)
    quality_rows = [
        ["Financial year", row["financial_year"] if pd.notna(row["financial_year"]) else "N/A", "Reference period"],
        ["ROE", _fmt(row["roe"], suffix="%"), _fmt(row["roe_risk"])],
        ["ROCE", _fmt(row["roce"], suffix="%"), _fmt(row["roce_risk"])],
        ["Debt/Equity", _fmt(row["debt_equity"]), _fmt(row["debt_equity_risk"])],
        [
            "Interest Cover",
            "Not Applicable" if row["is_financial"] or row["is_etf"] else _fmt(row["interest_cover"], suffix="x"),
            "N/A" if row["is_financial"] or row["is_etf"] else _fmt(row["interest_cover_risk"]),
        ],
    ]
    _add_table(document, ["Metric", "Raw Value", "Normalized Risk (0-100)"], quality_rows)
    document.add_paragraph(f"Status: {row['quality_status']}.")
    if pd.notna(row["quality_score"]):
        applicable = [
            row["roe_risk"],
            row["roce_risk"],
            row["debt_equity_risk"],
        ]
        if not row["is_financial"]:
            applicable.append(row["interest_cover_risk"])
        document.add_paragraph(
            "Stock Quality Risk = average("
            + ", ".join(f"{float(value):.2f}" for value in applicable)
            + f") = {row['quality_score']:.2f}."
        )
    else:
        document.add_paragraph("Stock Quality Risk = Not Applicable / not included.")

    document.add_heading("Liquidity calculation", level=2)
    liquidity_tiers = config["liquidity"]
    fallback_tier_name = config["liquidity_fallback_tier"]
    liquidity_tier = liquidity_tiers.get(
        row["cap_category"], liquidity_tiers[fallback_tier_name]
    )
    _add_table(
        document,
        ["Observations", "Average Volume ('000)", "Volume Risk", "Average Turnover (Cr)", "Turnover Risk", "Stock Liquidity Risk"],
        [[
            _fmt(row["liquidity_observations"], 0),
            _fmt(row["average_volume"]),
            _fmt(row["volume_risk"]),
            _fmt(row["average_turnover"]),
            _fmt(row["turnover_risk"]),
            _fmt(row["liquidity_score"]),
        ]],
    )
    tier_label = (
        row["cap_category"] if row["cap_category"] in liquidity_tiers else f"{fallback_tier_name} (fallback)"
    )
    document.add_paragraph(
        f"Volume and turnover use logarithmic inverse normalization between the "
        f"{tier_label} band: Volume {liquidity_tier['volume_low_thousand']}-"
        f"{liquidity_tier['volume_high_thousand']} ('000), Turnover "
        f"{liquidity_tier['turnover_low_crore']}-{liquidity_tier['turnover_high_crore']} (Cr). "
        "Stock Liquidity Risk = (Volume Risk + Turnover Risk) / 2."
    )

    document.add_heading("Volatility calculation", level=2)
    _add_table(
        document,
        ["Return Observations", "Annual Volatility", "Risk", "Maximum Drawdown", "Risk", "Downside Volatility", "Risk", "Stock Volatility Risk"],
        [[
            _fmt(row["return_observations"], 0),
            _fmt(row["annual_volatility"] * 100 if pd.notna(row["annual_volatility"]) else np.nan, suffix="%"),
            _fmt(row["annual_volatility_risk"]),
            _fmt(row["maximum_drawdown"] * 100 if pd.notna(row["maximum_drawdown"]) else np.nan, suffix="%"),
            _fmt(row["drawdown_risk"]),
            _fmt(row["downside_volatility"] * 100 if pd.notna(row["downside_volatility"]) else np.nan, suffix="%"),
            _fmt(row["downside_risk"]),
            _fmt(row["volatility_score"]),
        ]],
    )
    document.add_paragraph(
        "Annual Volatility Risk = Annual Volatility / 60% x 100; Drawdown Risk = "
        "Maximum Drawdown / 50% x 100; Downside Risk = Downside Volatility / 45% "
        "x 100. Each is capped at 100. Stock Volatility Risk is their average."
    )

    document.add_heading("Beta calculation", level=2)
    beta_benchmark = row["beta_benchmark"] if pd.notna(row.get("beta_benchmark")) else "N/A"
    _add_table(
        document,
        ["Benchmark Used", "Overlapping Observations", "Stock Beta", "Normalized Beta Risk"],
        [[
            beta_benchmark,
            _fmt(row["beta_observations"], 0),
            _fmt(row["beta"], 4),
            _fmt(row["beta_score"]),
        ]],
    )
    if pd.notna(row["beta"]):
        low = config["beta"]["low"]
        high = config["beta"]["high"]
        document.add_paragraph(
            f"Beta = Covariance(stock returns, {beta_benchmark} returns) / "
            f"Variance({beta_benchmark} returns) = {row['beta']:.4f}. Beta Risk = "
            f"clip(({row['beta']:.4f} - {low:.2f}) / ({high:.2f} - {low:.2f}) x 100) = "
            f"{row['beta_score']:.2f}. Benchmark chosen by market-cap category "
            f"({row['cap_category']})."
        )
    else:
        document.add_paragraph("Beta could not be calculated because coverage was insufficient.")


def generate() -> Path:
    portfolio, market, benchmark, config = _load_data()
    metrics = _calculate_stock_metrics(portfolio, market, benchmark, config)
    engine = build_engine(Settings.from_env().database_url)
    risk_input = metrics[
        ["security_name", "isin", "value", "sector", "nse_sector"]
    ].rename(columns={"sector": "ace_sector"})
    analysis = Riskometer(engine, ROOT / "riskometer_config.json").analyze(risk_input)

    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.color.rgb = RGBColor(15, 23, 42)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.color.rgb = RGBColor(29, 78, 216)
    _add_document_header(document)
    _add_methodology(document, metrics, analysis, config)
    for rank, (_, row) in enumerate(metrics.iterrows(), start=1):
        _add_stock_section(document, row, rank, metrics, config)

    document.add_page_break()
    document.add_heading("Final Reconciliation", level=1)
    document.add_paragraph(
        f"The individual and portfolio-level calculations reconcile to an overall "
        f"Risk-O-Meter score of {analysis['overall_score']:.2f}/100, classified as "
        f"{analysis['overall_level']} Risk."
    )
    document.add_paragraph(
        "This report is an analytical explanation of the configured Risk-O-Meter "
        "and is not investment advice."
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(generate())
