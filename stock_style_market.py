"""Momentum and Low-Volatility raw metrics, computed from the Accord
Data_NSE / Data_BSE monthly price files and the Data_Nifty_50.csv benchmark
series (Portfolio Analyzer Data folder).

NSE close is used as the primary price source; BSE is used only to fill in
ISINs that have little or no NSE coverage in the loaded window.
"""

import glob
import os
import re
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from config import STYLE_BSE_PRICE_DIR, STYLE_NIFTY_FILE, STYLE_NSE_PRICE_DIR

MONTHS_LOADED = 38  # 3 low-volatility years + buffer for 200DMA / lookback
MIN_NSE_OBSERVATIONS = 40  # below this, an ISIN is topped up from BSE
CACHE_DIR = Path(".style_cache")


def _month_key(path):
    name = os.path.splitext(os.path.basename(path))[0]
    return int(name[4:8]), int(name[2:4])


def _sorted_month_files(directory):
    files = glob.glob(os.path.join(directory, "*.xlsx"))
    files = [f for f in files if re.fullmatch(r"\d{8}", os.path.splitext(os.path.basename(f))[0])]
    files.sort(key=_month_key)
    return files


def _parse_month_file(path, isin_idx, date_idx, close_idx):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        isin, date, close = row[isin_idx], row[date_idx], row[close_idx]
        if isin is None or date is None or close is None:
            continue
        rows.append((str(isin).strip().upper(), pd.Timestamp(date).normalize(), float(close)))
    wb.close()
    return rows


def _cache_signature(files):
    return "|".join(f"{f}:{os.path.getmtime(f)}" for f in files)


def _load_price_history(directory, isin_idx, date_idx, close_idx, cache_name):
    files = _sorted_month_files(directory)[-MONTHS_LOADED:]
    if not files:
        return pd.DataFrame()

    CACHE_DIR.mkdir(exist_ok=True)
    cache_file = CACHE_DIR / f"{cache_name}.parquet"
    sig_file = CACHE_DIR / f"{cache_name}.sig"
    signature = _cache_signature(files)

    if cache_file.exists() and sig_file.exists() and sig_file.read_text(encoding="utf-8") == signature:
        return pd.read_parquet(cache_file)

    rows = []
    for path in files:
        rows.extend(_parse_month_file(path, isin_idx, date_idx, close_idx))

    long_df = pd.DataFrame(rows, columns=["isin", "date", "close"])
    long_df = long_df.drop_duplicates(subset=["isin", "date"], keep="last")
    wide = long_df.pivot(index="date", columns="isin", values="close").sort_index()

    wide.to_parquet(cache_file)
    sig_file.write_text(signature, encoding="utf-8")
    return wide


def load_nse_price_history():
    return _load_price_history(STYLE_NSE_PRICE_DIR, 3, 5, 6, "nse_prices")


def load_bse_price_history():
    return _load_price_history(STYLE_BSE_PRICE_DIR, 3, 5, 6, "bse_prices")


def load_combined_price_history():
    """NSE close as primary series; BSE fills ISINs with thin NSE coverage."""
    nse = load_nse_price_history()
    bse = load_bse_price_history()

    nse_counts = nse.count()
    thin_isins = nse_counts[nse_counts < MIN_NSE_OBSERVATIONS].index
    missing_isins = bse.columns.difference(nse.columns)
    fill_isins = [isin for isin in thin_isins if isin in bse.columns] + list(missing_isins)

    overlap = [isin for isin in fill_isins if isin in nse.columns]
    new_only = [isin for isin in fill_isins if isin not in nse.columns]

    combined = nse.drop(columns=overlap)
    filled_overlap = pd.DataFrame(
        {isin: nse[isin].combine_first(bse[isin]) for isin in overlap}
    )
    combined = pd.concat([combined, filled_overlap, bse[new_only]], axis=1)

    return combined.sort_index(axis=0).sort_index(axis=1)


def load_nifty_series():
    df = pd.read_csv(STYLE_NIFTY_FILE)
    df["date"] = pd.to_datetime(df["Date"], format="%m/%d/%Y")
    df["close"] = df["Price"].astype(str).str.replace(",", "", regex=False).astype(float)
    series = df.set_index("date")["close"].sort_index()
    return series[~series.index.duplicated(keep="last")]


def _asof(series, target_date):
    valid = series.dropna()
    if valid.empty or target_date < valid.index.min():
        return np.nan
    return valid.asof(target_date)


def compute_momentum_low_vol_universe(config):
    prices = load_combined_price_history()
    nifty = load_nifty_series()

    if prices.empty:
        return pd.DataFrame()

    history = config["history"]
    dma_window = history["dma_window"]
    trading_days = history["trading_days_per_year"]
    low_vol_years = history["low_vol_years"]
    min_obs = history["minimum_return_observations"]

    latest_date = prices.index.max()
    nifty_now = _asof(nifty, latest_date)
    nifty_6m = _asof(nifty, latest_date - pd.DateOffset(months=6))
    nifty_12m = _asof(nifty, latest_date - pd.DateOffset(months=12))
    nifty_return_12m = (nifty_now / nifty_12m - 1) if nifty_12m else np.nan

    cutoff = latest_date - pd.DateOffset(years=low_vol_years)
    nifty_window = nifty[nifty.index >= cutoff]
    nifty_returns = nifty_window.pct_change().dropna()

    records = {}
    for isin in prices.columns:
        series = prices[isin].dropna()
        if series.empty:
            continue

        price_now = series.iloc[-1]
        price_6m = _asof(series, latest_date - pd.DateOffset(months=6))
        price_12m = _asof(series, latest_date - pd.DateOffset(months=12))

        return_6m = price_now / price_6m - 1 if price_6m else np.nan
        return_12m = price_now / price_12m - 1 if price_12m else np.nan
        relative_strength = (
            (1 + return_12m) / (1 + nifty_return_12m)
            if pd.notna(return_12m) and pd.notna(nifty_return_12m)
            else np.nan
        )

        if len(series) >= dma_window:
            dma_200 = series.tail(dma_window).mean()
            above_200dma = bool(price_now > dma_200)
        else:
            above_200dma = np.nan

        window = series[series.index >= cutoff]
        stock_returns = window.pct_change().dropna()

        beta = np.nan
        annualized_volatility = np.nan
        max_drawdown = np.nan
        if len(stock_returns) >= min_obs:
            annualized_volatility = float(stock_returns.std() * np.sqrt(trading_days))
            wealth = (1 + stock_returns).cumprod()
            max_drawdown = float(abs((wealth / wealth.cummax() - 1).min()))

            aligned = pd.concat(
                [stock_returns.rename("stock"), nifty_returns.rename("nifty")], axis=1
            ).dropna()
            if len(aligned) >= min_obs and aligned["nifty"].var() > 0:
                beta = float(aligned["stock"].cov(aligned["nifty"]) / aligned["nifty"].var())

        records[isin] = {
            "return_6m": return_6m,
            "return_12m": return_12m,
            "relative_strength": relative_strength,
            "above_200dma": above_200dma,
            "beta": beta,
            "annualized_volatility": annualized_volatility,
            "max_drawdown": max_drawdown,
        }

    result = pd.DataFrame.from_dict(records, orient="index")
    result.index.name = "isin"
    result.attrs["as_of_date"] = latest_date
    return result
